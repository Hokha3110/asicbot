import os
import re
from typing import List, Dict, Any

class DocumentParserService:
    @staticmethod
    def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
        """
        Parses PDF using PyMuPDF (fitz) and extracts text per page with metadata.
        """
        pages_content = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text").strip()
                if text:
                    pages_content.append({
                        "page_number": page_num + 1,
                        "text": text,
                        "char_count": len(text)
                    })
            doc.close()
        except ImportError:
            # Fallback if fitz is missing
            pages_content = DocumentParserService._fallback_text_extract(file_path)
        except Exception as e:
            print(f"Error parsing PDF {file_path}: {e}")
            pages_content = DocumentParserService._fallback_text_extract(file_path)
            
        return pages_content

    @staticmethod
    def parse_docx(file_path: str) -> List[Dict[str, Any]]:
        """
        Parses DOCX using python-docx and extracts text sections.
        """
        sections_content = []
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))
                        
            joined_text = "\n\n".join(full_text)
            # Group into synthetic pages of ~2000 characters
            chunks = DocumentParserService._chunk_into_pages(joined_text)
            for idx, chunk in enumerate(chunks):
                sections_content.append({
                    "page_number": idx + 1,
                    "text": chunk,
                    "char_count": len(chunk)
                })
        except Exception as e:
            print(f"Error parsing DOCX {file_path}: {e}")
            sections_content = DocumentParserService._fallback_text_extract(file_path)
            
        return sections_content

    @staticmethod
    def parse_pptx(file_path: str) -> List[Dict[str, Any]]:
        """
        Parses PPTX using python-pptx and extracts text per slide.
        """
        slides_content = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for idx, slide in enumerate(prs.slides):
                slide_text_list = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_text_list.append(paragraph.text.strip())
                
                # Slide notes if any
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text.strip()
                    if notes:
                        slide_text_list.append(f"Notes: {notes}")
                        
                combined_slide_text = "\n".join(slide_text_list).strip()
                if combined_slide_text:
                    slides_content.append({
                        "page_number": idx + 1,
                        "text": combined_slide_text,
                        "char_count": len(combined_slide_text)
                    })
        except Exception as e:
            print(f"Error parsing PPTX {file_path}: {e}")
            slides_content = DocumentParserService._fallback_text_extract(file_path)
            
        return slides_content

    @staticmethod
    def chunk_document(pages_content: List[Dict[str, Any]], doc_name: str, vendor: str, category: str, chunk_size: int = 1000, overlap: int = 150) -> List[Dict[str, Any]]:
        """
        Splits pages into semantic chunks with full Presales metadata.
        """
        chunks = []
        for item in pages_content:
            page_num = item["page_number"]
            page_text = item["text"]
            
            if len(page_text) <= chunk_size:
                chunks.append({
                    "chunk_id": f"{doc_name}_p{page_num}_c0",
                    "text": page_text,
                    "metadata": {
                        "document_name": doc_name,
                        "vendor": vendor,
                        "solution_category": category,
                        "page_number": page_num
                    }
                })
            else:
                start = 0
                c_idx = 0
                while start < len(page_text):
                    end = min(start + chunk_size, len(page_text))
                    # Try to break at a paragraph or sentence boundary
                    if end < len(page_text):
                        boundary = page_text.rfind("\n", start, end)
                        if boundary == -1 or boundary < start + (chunk_size // 2):
                            boundary = page_text.rfind(". ", start, end)
                        if boundary != -1 and boundary > start + (chunk_size // 2):
                            end = boundary + 1
                            
                    sub_chunk = page_text[start:end].strip()
                    if sub_chunk:
                        chunks.append({
                            "chunk_id": f"{doc_name}_p{page_num}_c{c_idx}",
                            "text": sub_chunk,
                            "metadata": {
                                "document_name": doc_name,
                                "vendor": vendor,
                                "solution_category": category,
                                "page_number": page_num
                            }
                        })
                        c_idx += 1
                    start = end - overlap if end < len(page_text) else len(page_text)

        return chunks

    @staticmethod
    def detect_vendor_and_category(text: str, filename: str) -> tuple[str, str]:
        """
        Heuristic detector for cybersecurity vendor and category from document text and filename.
        """
        lower = (filename + " " + text[:2000]).lower()
        
        # Vendor detection
        vendor = "Cybersecurity Vendor"
        if "picus" in lower:
            vendor = "Picus Security"
        elif "graylog" in lower:
            vendor = "Graylog Inc."
        elif "trend micro" in lower or "xdr" in lower:
            vendor = "Trend Micro"
        elif "forcepoint" in lower:
            vendor = "Forcepoint"
        elif "netwrix" in lower:
            vendor = "Netwrix Corporation"
        elif "certus" in lower:
            vendor = "Certus Cybersecurity"
        elif "crowdstrike" in lower:
            vendor = "CrowdStrike"
        elif "palo alto" in lower:
            vendor = "Palo Alto Networks"
        elif "fortinet" in lower:
            vendor = "Fortinet"
        elif "sentinelone" in lower:
            vendor = "SentinelOne"

        # Category detection
        category = "General Security"
        if any(w in lower for w in ["dspm", "data security posture", "data classification", "sensitive data"]):
            category = "DSPM (Data Security Posture Management)"
        elif any(w in lower for w in ["bas", "breach and attack simulation", "mitre", "ctem", "continuous threat"]):
            category = "BAS (Breach & Attack Simulation)"
        elif any(w in lower for w in ["siem", "soc", "log management", "security operations"]):
            category = "SOC / SIEM / Log Management"
        elif any(w in lower for w in ["xdr", "edr", "endpoint protection", "extended detection"]):
            category = "XDR / EDR Endpoint Security"
        elif any(w in lower for w in ["dlp", "data loss prevention", "leakage"]):
            category = "DLP (Data Loss Prevention)"
        elif any(w in lower for w in ["pam", "privileged access", "vault", "credential"]):
            category = "PAM (Privileged Access Management)"
        elif any(w in lower for w in ["access analyzer", "identity", "iam", "active directory", "permissions"]):
            category = "Identity & Access Governance"
        elif any(w in lower for w in ["devsecops", "sast", "dast", "sca"]):
            category = "DevSecOps & Application Security"

        return vendor, category

    @staticmethod
    def _chunk_into_pages(text: str, page_size: int = 2000) -> List[str]:
        chunks = []
        for i in range(0, len(text), page_size):
            chunks.append(text[i:i + page_size])
        return chunks if chunks else [text]

    @staticmethod
    def _fallback_text_extract(file_path: str) -> List[Dict[str, Any]]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            return [{"page_number": 1, "text": content[:5000], "char_count": len(content)}]
        except Exception:
            return [{"page_number": 1, "text": "Preview text unavailable for this file.", "char_count": 0}]
